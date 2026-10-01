from datetime import datetime

import pytest
from openpyxl import Workbook

from app.nfl_predictor.report import calculated, parse_workbook, settle


@pytest.mark.parametrize(
    "market,side,line,home,away,result",
    [
        ("ML", "H", None, 24, 21, "win"),
        ("ML", "A", None, 24, 21, "loss"),
        ("ATS", "H", -3, 24, 21, "push"),
        ("ATS", "A", 3.5, 24, 21, "win"),
        ("TOTAL", "U", 45, 24, 21, "push"),
        ("CONTRA_TOTAL", "U", 46, 24, 21, "win"),
        ("ATS", "H", None, 24, 21, "pending"),
        ("ML", "H", None, None, None, "pending"),
    ],
)
def test_settlement(market, side, line, home, away, result):
    assert settle(market, side, line, home, away) == result


def test_kelly_accounts_for_push_and_negative_edge():
    result = calculated(0.6, 0.1, 2, 100, "push")
    assert result["ev"] == pytest.approx(0.3)
    assert result["stake"] == 33.33
    assert result["profit"] == 0
    assert calculated(0.4, 0, 2, 100, "loss")["stake"] == 0
    assert calculated(0.6, 0, None, 100, "win")["profit"] is None


def test_import_preserves_original_and_recalculates_contras(tmp_path):
    book = Workbook()
    ws = book.active
    ws.title = "NFL2025"
    ws["BG30"] = "Kelly"
    ws["C30"] = "Game_ID"
    master = book.create_sheet("NFL22232425")
    master.append(["headers"])
    master.append(
        [
            1,
            2025,
            17,
            "REG",
            12,
            25,
            2025,
            "KC",
            "DEN",
            24,
            21,
            -150,
            130,
            -3,
            3,
            45,
            -110,
            -105,
            -120,
            110,
        ]
    )
    master.append([2, 2025])
    preds = book.create_sheet("NFL_25P")
    preds.append(["headers"])
    preds.append(
        [1, datetime(2025, 12, 25), 17, "KC", "DEN", "H", 0.6, "H", 0.55, 0.1, "O", 0.52, 0.08]
    )
    preds.append([2])
    for cell in ["BG28", "BP28", "BX28", "CF28", "CN28"]:
        ws[cell] = 100
    ws["D31"] = datetime(2025, 12, 25)
    ws["E31"] = "KC"
    ws["F31"] = "DEN"
    ws["C31"] = 1
    ws["C32"] = 2
    for col, value in {
        "I": "H",
        "J": "H",
        "K": "O",
        "L": "A",
        "M": "U",
        "AN": "L",
        "BR": -42,
        "BP": 42,
    }.items():
        ws[f"{col}31"] = value
    path = tmp_path / "source.xlsx"
    book.save(path)
    rows, meta = parse_workbook(path)
    assert len(rows) == 5
    assert meta["skipped_rows"] == 1
    markets = {r["market"]: r for r in rows}
    assert markets["ATS"]["original"]["profit"] == -42
    assert markets["ATS"]["corrected"]["outcome"] == "push"
    assert markets["CONTRA_ATS"]["corrected"]["probability"] == pytest.approx(0.35)
    assert markets["CONTRA_TOTAL"]["corrected"]["probability"] == pytest.approx(0.4)
    assert markets["CONTRA_TOTAL"]["corrected"]["odds"] == 2.1


@pytest.mark.parametrize("role,active", [("user", True), ("admin", True), ("master_admin", False)])
def test_report_restricted(client, monkeypatch, role, active):
    from test_nfl_predictor_online import set_profile_role

    from app.api.v1.routes import nfl_predictor

    set_profile_role(role, active)

    def forbidden():
        pytest.fail("Unauthorized database access")

    monkeypatch.setattr(nfl_predictor, "connection", forbidden)
    assert client.get("/api/v1/nfl-predictor/report").status_code == 403


def test_comparison_deduplicates_excel_and_uses_common_games():
    from app.nfl_predictor.report import comparison

    common = dict(date="2025-12-25", home="KC", away="DEN", stake=10, profit=10, outcome="win")
    absent = dict(common, home="BUF")
    result = comparison([common, absent], [common, dict(common)])
    assert result["games"] == 1
    assert result["duplicate_excel_rows"] == 1
    assert result["engine"]["profit"] == result["original"]["profit"] == 10


def test_engine_predictions_do_not_change_when_target_scores_change():
    import numpy as np
    import pandas as pd

    from app.nfl_predictor import data, pipeline
    from app.nfl_predictor.backtest import selections

    random = np.random.default_rng(17)
    records = []
    for i in range(96):
        records.append(
            dict(
                Game_ID=str(i),
                date=pd.Timestamp("2020-09-01") + pd.Timedelta(days=(i // 4) * 7),
                season=2020,
                week=i // 4 + 1,
                home_team=["KC", "BUF", "DEN", "SF"][i % 4],
                away_team=["BAL", "MIA", "LV", "SEA"][i % 4],
                home_points=float(random.integers(7, 40)),
                away_points=float(random.integers(7, 40)),
                home_ml=-110.0,
                away_ml=-110.0,
                spread_close=-3.0,
                total_close=45.0,
                home_spread_odds=-110.0,
                away_spread_odds=-110.0,
                over_odds=-110.0,
                under_odds=-110.0,
            )
        )
    games = pd.DataFrame(records)
    targets = [(2020, w) for w in range(1, 25)]
    original = pipeline.run(data.prepare(games, None, targets), [(2020, 24)])
    games.loc[games.week == 24, ["home_points", "away_points"]] = [99.0, 0.0]
    altered = pipeline.run(data.prepare(games, None, targets), [(2020, 24)])
    columns = ["p_home_ensemble", "p_home_cover", "p_over", "lambda_home", "lambda_away"]
    np.testing.assert_allclose(original[columns], altered[columns], rtol=0, atol=1e-12)
    entries = selections(
        original.iloc[0], {k: 100 for k in ["ML", "ATS", "TOTAL", "CONTRA_ATS", "CONTRA_TOTAL"]}
    )
    assert len(entries) == 5
    assert all(r["training_cutoff"] < "2021-02-10" for r in entries)
    markets = {r["market"]: r for r in entries}
    assert markets["ATS"]["selection"] != markets["CONTRA_ATS"]["selection"]
    assert markets["TOTAL"]["selection"] != markets["CONTRA_TOTAL"]["selection"]


def test_comparison_explains_stakes_pushes_and_no_entries():
    from app.nfl_predictor.report import comparison

    rows = [
        dict(
            date="2025-12-25", home=str(i), away="DEN", stake=stake, profit=profit, outcome=outcome
        )
        for i, (stake, profit, outcome) in enumerate(
            [(20, 10, "win"), (10, -10, "loss"), (30, 0, "push"), (0, 0, "loss")]
        )
    ]
    result = comparison(rows, rows)
    stats = result["engine"]
    assert result["games"] == 4
    assert stats["staked"] == 60
    assert stats["wins"] + stats["losses"] + stats["pushes"] == stats["entries"] == 3
    assert stats["no_entry"] == 1
    assert stats["roi"] == 0
