"""Import and inspect the user's Kelly workbook without recording real wagers."""

import argparse
import hashlib
import math
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from psycopg.types.json import Jsonb

from .data import team_name
from .pipeline import decimal_odds
from .store import connection

MARKETS = {
    "ML": ("Moneyline", "I", "AF", "AG", "BF", "BG", "BH", "BJ", "AH", "BG28"),
    "ATS": ("Spread", "J", "AL", "AM", "BN", "BO", "BP", "BR", "AN", "BP28"),
    "TOTAL": ("Totales", "K", "AQ", "AR", "BV", "BW", "BX", "BZ", "AS", "BX28"),
    "CONTRA_ATS": ("Contra spread", "L", "AV", "AW", "CD", "CE", "CF", "CH", "AX", "CF28"),
    "CONTRA_TOTAL": ("Contra total", "M", "BA", "BB", "CL", "CM", "CN", "CP", "BC", "CN28"),
}


def number(value):
    if isinstance(value, bool):
        return None
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (TypeError, ValueError):
        return None


def settle(market, side, line, home, away):
    if home is None or away is None:
        return "pending"
    if market == "ML":
        margin = home - away if side == "H" else away - home
    elif market in ("ATS", "CONTRA_ATS"):
        if line is None:
            return "pending"
        margin = (home - away if side == "H" else away - home) + line
    else:
        if line is None:
            return "pending"
        margin = home + away - line if side == "O" else line - home - away
    return "win" if margin > 0 else "loss" if margin < 0 else "push"


def calculated(probability, push, odds, base, outcome):
    if (
        probability is None
        or push is None
        or odds is None
        or base is None
        or odds <= 1
        or not 0 <= probability <= 1 - push + 1e-8
    ):
        return dict(
            probability=probability,
            push=push,
            odds=odds,
            ev=None,
            kelly=None,
            stake=None,
            profit=None,
            outcome=outcome,
        )
    loss = max(0.0, 1 - probability - push)
    ev = probability * (odds - 1) - loss
    fraction = max(0.0, min(1.0, ev / ((odds - 1) * (1 - push)))) if push < 1 else 0.0
    stake = round(fraction * base, 2)
    profit = (
        round(stake * (odds - 1), 2)
        if outcome == "win"
        else -stake
        if outcome == "loss"
        else 0.0
        if outcome == "push"
        else None
    )
    return dict(
        probability=probability,
        push=push,
        odds=odds,
        ev=ev,
        kelly=fraction,
        stake=stake,
        profit=profit,
        outcome=outcome,
    )


def parse_workbook(path):
    book = load_workbook(path, read_only=False, data_only=True)
    try:
        ws = book["NFL2025"]
        master = book["NFL22232425"]
        predictions = book["NFL_25P"]
        if ws["BG30"].value != "Kelly" or ws["C30"].value != "Game_ID":
            raise ValueError("La estructura del reporte cambió; no se importó.")
        games = {}
        for r in master.iter_rows(min_row=2, values_only=True):
            r = list(r)
            if r[6] is None and r[4] in (1, 2, 9, 10, 11, 12) and r[1] is not None:
                r[6] = int(str(r[1]).replace("NFL", "")) + int(r[4] <= 2)
            if all(r[c] is not None for c in (4, 5, 6, 7, 8)):
                key = (
                    datetime(int(r[6]), int(r[4]), int(r[5])).date(),
                    team_name(r[7]),
                    team_name(r[8]),
                )
                if key in games:
                    raise ValueError("Partido duplicado por fecha/equipos en el master.")
                games[key] = r
        probs = {}
        for r in predictions.iter_rows(min_row=2, values_only=True):
            if isinstance(r[1], datetime):
                probs[(r[1].date(), team_name(r[3]), team_name(r[4]))] = r
        bases = {key: number(ws[spec[-1]].value) for key, spec in MARKETS.items()}
        rows = []
        skipped = 0
        for idx in range(31, ws.max_row + 1):
            game_id = number(ws[f"C{idx}"].value)
            if game_id is None:
                skipped += 1
                continue
            report_date = ws[f"D{idx}"].value
            if not isinstance(report_date, datetime):
                skipped += 1
                continue
            key = (
                report_date.date(),
                team_name(ws[f"E{idx}"].value),
                team_name(ws[f"F{idx}"].value),
            )
            g = games.get(key)
            pred = probs.get(key)
            if not g or not pred:
                skipped += 1
                continue
            original_id = f"excel-row-{idx}"
            season = int(str(g[1]).replace("NFL", ""))
            week = int(g[2])
            date = datetime(int(g[6]), int(g[4]), int(g[5])).date().isoformat()
            home, away = team_name(g[7]), team_name(g[8])
            hp, ap = number(g[9]), number(g[10])
            for key, spec in MARKETS.items():
                (
                    label,
                    selection_col,
                    odds_col,
                    p_col,
                    ev_col,
                    kelly_col,
                    stake_col,
                    profit_col,
                    result_col,
                    _,
                ) = spec
                get = lambda col: ws[f"{col}{idx}"].value
                side = get(selection_col)
                if (
                    side not in ("H", "A")
                    if key in ("ML", "ATS", "CONTRA_ATS")
                    else side not in ("O", "U")
                ):
                    continue
                line = (
                    None
                    if key == "ML"
                    else number(g[13])
                    if key in ("ATS", "CONTRA_ATS")
                    else number(g[15])
                )
                if key in ("ATS", "CONTRA_ATS") and side == "A" and line is not None:
                    line = -line
                selection = (
                    home
                    if side == "H"
                    else away
                    if side == "A"
                    else "Over"
                    if side == "O"
                    else "Under"
                )
                original = dict(
                    probability=number(get(p_col)),
                    push=None,
                    odds=number(get(odds_col)),
                    ev=number(get(ev_col)),
                    kelly=number(get(kelly_col)),
                    stake=number(get(stake_col)),
                    profit=number(get(profit_col)),
                    outcome={"W": "win", "L": "loss", "P": "push"}.get(
                        str(get(result_col)), "pending"
                    ),
                )
                if key == "ML":
                    p = number(pred[6])
                    push = 0.0
                    price = g[11 if side == "H" else 12]
                elif key in ("ATS", "CONTRA_ATS"):
                    p = number(pred[8])
                    push = number(pred[9])
                    price = g[16 if side == "H" else 17]
                    if key == "CONTRA_ATS" and p is not None and push is not None:
                        p = 1 - p - push
                else:
                    p = number(pred[11])
                    push = number(pred[12])
                    price = g[18 if side == "O" else 19]
                    if key == "CONTRA_TOTAL" and p is not None and push is not None:
                        p = 1 - p - push
                odds = decimal_odds(number(price))
                odds = None if math.isnan(odds) else odds
                corrected = calculated(p, push, odds, bases[key], settle(key, side, line, hp, ap))
                rows.append(
                    dict(
                        game_id=original_id,
                        season=season,
                        week=week,
                        date=date,
                        home=home,
                        away=away,
                        home_score=hp,
                        away_score=ap,
                        market=key,
                        selection=selection,
                        line=line,
                        base=bases[key],
                        original=original,
                        corrected=corrected,
                    )
                )
        return rows, dict(
            filename=Path(path).name,
            sha256=hashlib.sha256(Path(path).read_bytes()).hexdigest(),
            rows=len(rows),
            skipped_rows=skipped,
            bases=bases,
            kind="historical_simulation",
        )
    finally:
        book.close()


def save_report(store, rows, metadata, report_id="muchadata2"):
    if not rows:
        raise ValueError("No hay filas válidas para importar.")
    with store.conn.transaction():
        store.conn.execute(
            "insert into public.nfp_kelly_reports(id,metadata) values(%s,%s) on conflict(id) do update set metadata=excluded.metadata,imported_at=now()",
            (report_id, Jsonb(metadata)),
        )
        store.conn.execute("delete from public.nfp_kelly_rows where report_id=%s", (report_id,))
        with store.conn.cursor() as cur:
            cur.executemany(
                "insert into public.nfp_kelly_rows(report_id,game_id,season,week,market,payload) values(%s,%s,%s,%s,%s,%s)",
                [
                    (report_id, r["game_id"], r["season"], r["week"], r["market"], Jsonb(r))
                    for r in rows
                ],
            )
    return metadata


def import_report(store, path):
    return save_report(store, *parse_workbook(path))


def comparison(engine, original):
    def key(row):
        return row["date"], row["home"], row["away"]

    baseline = {}
    for row in original:
        baseline.setdefault(key(row), row)
    pairs = [
        (r, baseline[key(r)])
        for r in engine
        if key(r) in baseline
        and r["stake"] is not None
        and baseline[key(r)]["stake"] is not None
        and r["profit"] is not None
        and baseline[key(r)]["profit"] is not None
    ]

    def stats(rows):
        active = [r for r in rows if r["stake"] > 0 and r["outcome"] in ("win", "loss", "push")]
        stake = sum(r["stake"] for r in active)
        profit = sum(r["profit"] for r in active)
        return dict(
            profit=round(profit, 2),
            staked=round(stake, 2),
            pushes=sum(r["outcome"] == "push" for r in active),
            no_entry=sum(r["stake"] == 0 for r in rows),
            roi=profit / stake if stake else None,
            wins=sum(r["outcome"] == "win" for r in active),
            losses=sum(r["outcome"] == "loss" for r in active),
            entries=len(active),
        )

    return dict(
        games=len(pairs),
        engine=stats([p[0] for p in pairs]),
        original=stats([p[1] for p in pairs]),
        duplicate_excel_rows=len(original) - len(baseline),
    )


def read_report(store, market="ML", season=None, week=None, mode="original"):
    report_id = "engine" if mode == "engine" else "muchadata2"
    payload_mode = "corrected" if mode == "engine" else mode
    metadata = store.conn.execute(
        "select metadata,imported_at from public.nfp_kelly_reports where id=%s", (report_id,)
    ).fetchone()
    params = [report_id, market]
    clauses = ["report_id=%s", "market=%s"]
    if season is not None:
        clauses.append("season=%s")
        params.append(season)
    if week is not None:
        clauses.append("week=%s")
        params.append(week)
    fetched = store.conn.execute(
        "select payload from public.nfp_kelly_rows where " + " and ".join(clauses), params
    )
    rows = []
    for item in fetched:
        raw = item["payload"]
        r = {k: v for k, v in raw.items() if k not in ("original", "corrected")}
        r.update(raw[payload_mode])
        rows.append(r)
    rows.sort(key=lambda r: (r["date"], r["game_id"]))
    accumulated = 0.0
    wins = losses = pushes = 0
    staked = 0.0
    for row in rows:
        amount = row["stake"]
        profit = row["profit"]
        settled = (
            row["outcome"] in ("win", "loss", "push")
            and amount is not None
            and amount > 0
            and profit is not None
        )
        if settled:
            accumulated += profit
            staked += amount
            wins += row["outcome"] == "win"
            losses += row["outcome"] == "loss"
            pushes += row["outcome"] == "push"
        row["cumulative"] = round(accumulated, 2)
    periods = list(
        store.conn.execute(
            "select distinct season,week from public.nfp_kelly_rows where report_id=%s order by season desc,week",
            (report_id,),
        )
    )
    return dict(
        rows=rows,
        periods=periods,
        source=metadata,
        mode=mode,
        comparison=comparison(rows, read_report(store, market, season, week, "original")["rows"])
        if mode == "engine"
        else None,
        summary=dict(
            profit=round(accumulated, 2),
            staked=round(staked, 2),
            wins=wins,
            losses=losses,
            pushes=pushes,
            roi=accumulated / staked if staked else None,
            win_rate=wins / (wins + losses) if wins + losses else None,
        ),
    )


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    with connection() as store:
        store.conn.execute(Path(__file__).with_name("report_schema.sql").read_text())
        print(import_report(store, args.workbook))


if __name__ == "__main__":
    main()
